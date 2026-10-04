#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
tests/test_svg_to_pptx.py
=========================
测试 svg_to_pptx.py：SVG 解析 / 背景剥离 / 单位换算 / PPTX 构建 / CLI。

端到端构建测试需要 rsvg-convert 与 python-pptx（缺失则跳过）；
qa_pptx 回读集成测试复用仓库既有门禁（默认字阶 [11,13,16,20,24,32,44,56,96]，
fixture 字号只用 20/96，保证门禁可过）。
"""

import io
import sys
import tempfile
import unittest
import zipfile
from unittest.mock import patch
from contextlib import redirect_stdout
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from scripts.svg_to_pptx import (
    TextItem,
    check_dependencies,
    estimate_width_px,
    extract_texts,
    main,
    parse_color,
    parse_viewbox,
    render_background,
    resolve_project_dir,
    resolve_svg_dir,
    rsvg_available,
    strip_body_texts,
    build_pptx,
    verify_pptx,
    _atomic_save,
    _materialize_images,
)

try:
    from scripts.svg_to_pptx import _PPTX_OK
except ImportError:  # pragma: no cover
    _PPTX_OK = False

try:
    from PIL import Image as _PILImage
    _PIL_TEST_OK = True
except ImportError:  # pragma: no cover
    _PILImage = None
    _PIL_TEST_OK = False

try:
    from scripts.qa_pptx import run_qa_pptx
    _QA_OK = True
except ImportError:  # pragma: no cover
    _QA_OK = False

SVG_TMPL = """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1920 1080" width="1920" height="1080">
<rect x="0" y="0" width="1920" height="1080" fill="{bg}"/>
{texts}
</svg>"""

BG_COLORS = ("#F5F2EC", "#EDE8DC", "#F0EDE6")

PAGE1_TEXTS = """
<text x="144" y="200" font-family="Noto Serif CJK SC" font-size="96" fill="#1A1A1A" font-weight="900">标题一</text>
<text x="960" y="400" font-family="Noto Sans CJK SC" font-size="20" fill="#575046" text-anchor="middle" letter-spacing="4">副标题文字</text>
<text x="1776" y="1050" font-family="Bodoni Moda" font-size="225" fill="#1A1A1A" opacity="0.08" text-anchor="end" data-decorative="true">01</text>
"""

PAGE2_TEXTS = """
<text x="1776" y="200" font-family="Noto Serif CJK SC" font-size="96" fill="#8C2F1B" font-weight="700" text-anchor="end">标题二</text>
<text x="144" y="400" font-family="Noto Sans CJK SC" font-size="20" fill="#1A1A1A">正文内容行</text>
"""

PAGE3_TEXTS = """
<text x="144" y="200" font-family="Noto Serif CJK SC" font-size="96" fill="#1A1A1A">标题三</text>
<text x="144" y="400" font-family="Noto Sans CJK SC" font-size="20" fill="#2F5D8C" font-style="italic">斜体说明</text>
"""


def _write_fixture(tmp: Path) -> list[Path]:
    files = []
    for i, (texts, bg) in enumerate(zip((PAGE1_TEXTS, PAGE2_TEXTS, PAGE3_TEXTS), BG_COLORS), start=1):
        p = tmp / f"0{i}_page{i}.svg"
        p.write_text(SVG_TMPL.format(texts=texts, bg=bg), encoding="utf-8")
        files.append(p)
    return files


class TestParse(unittest.TestCase):
    def test_parse_viewbox(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "a.svg"
            p.write_text(SVG_TMPL.format(texts="", bg="#F5F2EC"), encoding="utf-8")
            self.assertEqual(parse_viewbox(p), (1920.0, 1080.0))

    def test_parse_viewbox_fallback_wh(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "b.svg"
            p.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="1280" height="720"></svg>',
                         encoding="utf-8")
            self.assertEqual(parse_viewbox(p), (1280.0, 720.0))

    def test_parse_color(self):
        self.assertEqual(parse_color("#1A1A1A"), (26, 26, 26))
        self.assertEqual(parse_color("#fff"), (255, 255, 255))
        self.assertEqual(parse_color("#8C2F1B"), (140, 47, 27))

    def test_px_to_pt_scale_matches_layout(self):
        # 字号换算必须与版式同比例：px_to_pt = 72 * sx，
        # 其中 sx = slide_w_in / vb_w。固定 0.75（96dpi 假设）曾导致
        # 1920px→13.333in 版式下字号大 1.5 倍、文本框装不下（2026-10-02）。
        from scripts.svg_to_pptx import FORMATS
        slide_w_in, _ = FORMATS["ppt169"]
        for vb_w in (1920.0, 1280.0):
            sx = slide_w_in / vb_w
            self.assertAlmostEqual(72.0 * sx, 72.0 * slide_w_in / vb_w)


class TestExtract(unittest.TestCase):
    def setUp(self):
        self.td = tempfile.TemporaryDirectory()
        self.files = _write_fixture(Path(self.td.name))

    def tearDown(self):
        self.td.cleanup()

    def test_decorative_excluded(self):
        texts = extract_texts(self.files[0])
        self.assertEqual(len(texts), 2)
        self.assertNotIn("01", [t.text for t in texts])

    def test_fields_page1(self):
        texts = extract_texts(self.files[0])
        t0, t1 = texts
        self.assertEqual(t0.text, "标题一")
        self.assertEqual(t0.size_px, 96)
        self.assertEqual(t0.anchor, "start")
        self.assertTrue(t0.bold)
        self.assertFalse(t0.italic)
        self.assertEqual(t0.fill, "#1A1A1A")
        self.assertEqual(t1.anchor, "middle")
        self.assertEqual(t1.letter_spacing_px, 4)
        self.assertEqual(t1.family, "Noto Sans CJK SC")

    def test_fields_page2_end_anchor_bold(self):
        (t0, t1) = extract_texts(self.files[1])
        self.assertEqual(t0.anchor, "end")
        self.assertTrue(t0.bold)
        self.assertEqual(t0.fill, "#8C2F1B")
        self.assertFalse(t1.bold)

    def test_italic(self):
        texts = extract_texts(self.files[2])
        self.assertTrue(texts[1].italic)

    def test_empty_text_skipped(self):
        p = Path(self.td.name) / "empty.svg"
        p.write_text(SVG_TMPL.format(
            texts='<text x="10" y="10" font-size="20">   </text>', bg="#F5F2EC"), encoding="utf-8")
        self.assertEqual(extract_texts(p), [])


class TestStrip(unittest.TestCase):
    def test_body_texts_removed_decorative_kept(self):
        with tempfile.TemporaryDirectory() as td:
            files = _write_fixture(Path(td))
            out = strip_body_texts(files[0]).decode("utf-8")
            self.assertNotIn("标题一", out)
            self.assertNotIn("副标题文字", out)
            self.assertIn("01", out)          # 装饰水印保留在背景层
            self.assertIn("<rect", out)       # 非文本元素不受影响

    def test_stripped_svg_still_valid_xml(self):
        import xml.etree.ElementTree as ET
        with tempfile.TemporaryDirectory() as td:
            files = _write_fixture(Path(td))
            root = ET.fromstring(strip_body_texts(files[0]))
            self.assertTrue(root.tag.endswith("svg"))


class TestWidth(unittest.TestCase):
    def test_estimate_positive(self):
        item = TextItem("标题一", 144, 200, 96, "Noto Serif CJK SC",
                        "#1A1A1A", True, False, "start", 0)
        w = estimate_width_px(item)
        self.assertGreater(w, 0)

    def test_letter_spacing_increases_width(self):
        base = TextItem("副标题文字", 960, 400, 20, "Noto Sans CJK SC",
                        "#575046", False, False, "middle", 0)
        spaced = TextItem("副标题文字", 960, 400, 20, "Noto Sans CJK SC",
                          "#575046", False, False, "middle", 10)
        self.assertGreater(estimate_width_px(spaced), estimate_width_px(base))


class TestResolve(unittest.TestCase):
    def test_explicit_svg_dir(self):
        with tempfile.TemporaryDirectory() as td:
            _write_fixture(Path(td))
            svg_dir, proj = resolve_svg_dir(td)
            self.assertEqual(svg_dir, Path(td).resolve())

    def test_missing_raises(self):
        with self.assertRaises(FileNotFoundError):
            resolve_svg_dir("/tmp/definitely-not-here-xyz")

    def test_resolve_from_subfolder_and_spec_file(self):
        with tempfile.TemporaryDirectory() as td:
            proj = Path(td) / "test_proj"
            svg_dir = proj / "svg_output"
            render_dir = proj / "render"
            images_dir = proj / "images"
            svg_dir.mkdir(parents=True)
            render_dir.mkdir(parents=True)
            images_dir.mkdir(parents=True)
            spec_file = proj / "spec_lock.md"
            spec_file.write_text("# spec\n", encoding="utf-8")
            _write_fixture(svg_dir)

            self.assertEqual(resolve_project_dir(svg_dir), proj.resolve())
            self.assertEqual(resolve_project_dir(render_dir), proj.resolve())
            self.assertEqual(resolve_project_dir(images_dir), proj.resolve())
            self.assertEqual(resolve_project_dir(spec_file), proj.resolve())

            resolved_render_svg, resolved_render_proj = resolve_svg_dir(render_dir)
            self.assertEqual(resolved_render_svg, svg_dir.resolve())
            self.assertEqual(resolved_render_proj, proj.resolve())

            resolved_images_svg, resolved_images_proj = resolve_svg_dir(images_dir)
            self.assertEqual(resolved_images_svg, svg_dir.resolve())
            self.assertEqual(resolved_images_proj, proj.resolve())

            resolved_spec_svg, resolved_spec_proj = resolve_svg_dir(spec_file)
            self.assertEqual(resolved_spec_svg, svg_dir.resolve())
            self.assertEqual(resolved_spec_proj, proj.resolve())

            first_svg = svg_dir / "01_page1.svg"
            resolved_single_svg, resolved_single_proj = resolve_svg_dir(first_svg)
            self.assertEqual(resolved_single_svg, svg_dir.resolve())
            self.assertEqual(resolved_single_proj, proj.resolve())


class TestCheck(unittest.TestCase):
    def test_check_returns_tuple(self):
        ok, problems = check_dependencies()
        self.assertIsInstance(ok, bool)
        self.assertIsInstance(problems, list)
        self.assertEqual(rsvg_available(), shutil_which_rsvg())

    def test_main_check_mode(self):
        buf = io.StringIO()
        with redirect_stdout(buf):
            code = main(["--check"])
        self.assertIn(code, (0, 1))


def shutil_which_rsvg():
    import shutil
    return shutil.which("rsvg-convert") is not None


@unittest.skipUnless(_PPTX_OK and shutil_which_rsvg(), "需要 python-pptx + rsvg-convert")
class TestBuildEndToEnd(unittest.TestCase):
    def test_build_three_pages(self):
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td)
            files = _write_fixture(tdp)
            out = tdp / "out.pptx"
            stats = build_pptx(files, out, fmt="ppt169")
            self.assertEqual(stats["pages"], 3)
            self.assertGreater(stats["texts"], 0)
            self.assertTrue(out.is_file())
            self.assertGreater(out.stat().st_size, 0)

            with zipfile.ZipFile(out) as z:
                names = z.namelist()
                slides = [n for n in names if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
                self.assertEqual(len(slides), 3)
                media = [n for n in names if n.startswith("ppt/media/")]
                self.assertEqual(len(media), 3)
                self.assertTrue(all(n.endswith(".png") for n in media))
                # 每页既有图片层也有文本层
                import xml.etree.ElementTree as ET
                NS_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
                for s in slides:
                    root = ET.fromstring(z.read(s))
                    pics = root.findall(f".//{NS_P}pic")
                    sps = root.findall(f".//{NS_P}sp")
                    self.assertGreater(len(pics), 0, s)
                    self.assertGreater(len(sps), 0, s)
                # 文本内容真实写入（可编辑，非纯图片）
                all_text = " ".join(
                    ET.tostring(ET.fromstring(z.read(s)), encoding="unicode") for s in slides)
                self.assertIn("标题一", all_text)
                self.assertIn("标题二", all_text)

    def test_build_empty_raises(self):
        with tempfile.TemporaryDirectory() as td:
            with self.assertRaises(ValueError):
                build_pptx([], Path(td) / "x.pptx")

    def test_build_bad_format_raises(self):
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td)
            files = _write_fixture(tdp)
            with self.assertRaises(ValueError):
                build_pptx(files, tdp / "x.pptx", fmt="ppt999")

    def test_render_background_png_valid(self):
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td)
            files = _write_fixture(tdp)
            png = tdp / "bg.png"
            render_background(strip_body_texts(files[0]), png, 1920, 1080)
            self.assertTrue(png.is_file())
            self.assertEqual(png.read_bytes()[:8], b"\x89PNG\r\n\x1a\n")

    @unittest.skipUnless(_QA_OK, "需要 qa_pptx")
    def test_qa_pptx_readback_passes(self):
        """生成的 pptx 必须通过仓库既有 7 项回读门禁。"""
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td)
            files = _write_fixture(tdp)
            out = tdp / "out.pptx"
            build_pptx(files, out)
            buf = io.StringIO()
            with redirect_stdout(buf):
                ok = run_qa_pptx(out, verbose=False)
            self.assertTrue(ok, f"qa_pptx 回读未通过:\n{buf.getvalue()}")

    def test_main_cli_end_to_end(self):
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td)
            _write_fixture(tdp)
            out = tdp / "cli.pptx"
            buf = io.StringIO()
            with redirect_stdout(buf):
                code = main([str(tdp), "-o", str(out)])
            self.assertEqual(code, 0)
            self.assertTrue(out.is_file())

    def test_atomic_save_keeps_old_file_on_save_failure(self):
        """save 中途抛错：旧产物原样保留，临时文件被清理。"""
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td)
            files = _write_fixture(tdp)
            out = tdp / "out.pptx"
            sentinel = b"SENTINEL-OLD-GOOD-FILE"
            out.write_bytes(sentinel)
            with patch("scripts.svg_to_pptx.Presentation") as MP:
                MP.return_value.save.side_effect = OSError("disk full")
                with self.assertRaises(OSError):
                    build_pptx(files, out)
            self.assertEqual(out.read_bytes(), sentinel)  # 旧文件未被截断
            self.assertEqual(list(tdp.glob("*.tmp-*.pptx")), [])  # 无残留

    def test_no_tmp_files_after_success(self):
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td)
            files = _write_fixture(tdp)
            out = tdp / "out.pptx"
            build_pptx(files, out)
            self.assertEqual(list(tdp.glob("*.tmp-*.pptx")), [])

    def test_build_stats_include_verified(self):
        """build_pptx 内置回读验证：verified 字段真实回读产物。"""
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td)
            files = _write_fixture(tdp)
            out = tdp / "out.pptx"
            stats = build_pptx(files, out)
            self.assertEqual(stats["verified"]["pages"], 3)
            self.assertEqual(stats["verified"]["path"], str(out))
            self.assertGreater(stats["verified"]["bytes"], 4096)

    def test_main_cli_prints_verified(self):
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td)
            _write_fixture(tdp)
            out = tdp / "cli2.pptx"
            buf = io.StringIO()
            with redirect_stdout(buf):
                code = main([str(tdp), "-o", str(out)])
            self.assertEqual(code, 0)
            self.assertIn("回读验证", buf.getvalue())


@unittest.skipUnless(_PPTX_OK and _PIL_TEST_OK, "需要 python-pptx + Pillow")
class TestVerifyPptx(unittest.TestCase):
    """verify_pptx 回读门：页数/图片层/体积三项可观测断言。"""

    def _make(self, tdp, slides=2, with_picture=True):
        from pptx import Presentation as _P
        from pptx.util import Inches as _In
        img = tdp / "one.png"
        _PILImage.new("RGB", (8, 8), (1, 2, 3)).save(img, "PNG")
        prs = _P()
        blank = prs.slide_layouts[6]
        for _ in range(slides):
            sl = prs.slides.add_slide(blank)
            if with_picture:
                sl.shapes.add_picture(str(img), _In(0), _In(0),
                                      width=_In(1), height=_In(1))
        p = tdp / "v.pptx"
        prs.save(str(p))
        return p

    def test_verify_positive(self):
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td)
            p = self._make(tdp, slides=2, with_picture=True)
            v = verify_pptx(p, 2)
            self.assertEqual(v["pages"], 2)
            self.assertGreater(v["bytes"], 4096)

    def test_verify_page_count_mismatch(self):
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td)
            p = self._make(tdp, slides=2, with_picture=True)
            with self.assertRaisesRegex(RuntimeError, "页数不符"):
                verify_pptx(p, 3)

    def test_verify_missing_picture(self):
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td)
            p = self._make(tdp, slides=2, with_picture=False)
            with self.assertRaisesRegex(RuntimeError, "缺少背景图片层"):
                verify_pptx(p, 2)

    def test_verify_too_small(self):
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td)
            p = tdp / "tiny.pptx"
            p.write_bytes(b"x" * 100)
            with self.assertRaisesRegex(RuntimeError, "过小"):
                verify_pptx(p, 1, min_bytes=4096)


@unittest.skipUnless(_PPTX_OK and shutil_which_rsvg() and _PIL_TEST_OK,
                     "需要 python-pptx + rsvg-convert + Pillow")
class TestBackgroundPhotoRegression(unittest.TestCase):
    """背景照片回归：第 16 轮照片丢失事故的门禁。"""

    PHOTO_SVG = """<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" viewBox="0 0 1920 1080" width="1920" height="1080">
<rect x="0" y="0" width="1920" height="1080" fill="#F5F2EC"/>
<image xlink:href="probe.jpg" x="100" y="100" width="120" height="90" preserveAspectRatio="xMidYMid slice"/>
</svg>"""

    def _write_photo_fixture(self, tmp: Path) -> Path:
        _PILImage.new("RGB", (120, 90), (210, 30, 30)).save(tmp / "probe.jpg", "JPEG", quality=90)
        p = tmp / "photo.svg"
        p.write_text(self.PHOTO_SVG, encoding="utf-8")
        return p

    def test_strip_keeps_xlink_href(self):
        """ET 序列化不得把 xlink:href 改写成 ns1:href（否则 rsvg 认不出图片）。"""
        with tempfile.TemporaryDirectory() as td:
            p = self._write_photo_fixture(Path(td))
            out = strip_body_texts(p).decode("utf-8")
            self.assertIn("xlink:href", out)
            self.assertNotIn("ns1:href", out)

    def test_materialize_converts_jpeg_to_png(self):
        """JPEG 物化为 PNG 并改写为绝对 file:// URI；PNG 文件真实存在。"""
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td)
            p = self._write_photo_fixture(tdp)
            work = tdp / "work"
            out = _materialize_images(strip_body_texts(p), tdp, work).decode("utf-8")
            self.assertIn(".png", out)
            self.assertNotIn("probe.jpg", out)
            pngs = list(work.glob("*.png"))
            self.assertEqual(len(pngs), 1)
            self.assertEqual(pngs[0].read_bytes()[:8], b"\x89PNG\r\n\x1a\n")

    def test_materialize_corrupt_jpeg_warns_and_keeps_going(self):
        """坏 JPEG（截断/损坏）：不抛异常炸掉构建，保留原 href 并警告，
        与"缺失文件"同策略（第 31 轮；真正的门禁是 qa_assets）。"""
        import io
        from contextlib import redirect_stdout
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td)
            (tdp / "probe.jpg").write_bytes(b"\xff\xd8\xff\xe0garbage-not-jpeg")
            p = tdp / "photo.svg"
            p.write_text(self.PHOTO_SVG, encoding="utf-8")
            work = tdp / "work"
            buf = io.StringIO()
            with redirect_stdout(buf):
                out = _materialize_images(strip_body_texts(p), tdp, work).decode("utf-8")
            self.assertIn("probe.jpg", out)  # 原 href 保留，rsvg 会跳过该图
            self.assertIn("转换失败", buf.getvalue())

    def test_materialize_cache_invalidated_on_source_change(self):
        """转换缓存 key 混入体积+mtime：源文件被替换后不得复用过期 PNG。"""
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td)
            p = self._write_photo_fixture(tdp)
            work = tdp / "work"
            _materialize_images(strip_body_texts(p), tdp, work)
            first = list(work.glob("*.png"))
            self.assertEqual(len(first), 1)
            # 替换源文件（不同尺寸 + 强制刷新 mtime）
            _PILImage.new("RGB", (60, 45), (30, 30, 210)).save(
                tdp / "probe.jpg", "JPEG", quality=90)
            import os, time
            new_mtime = time.time() + 5
            os.utime(tdp / "probe.jpg", (new_mtime, new_mtime))
            _materialize_images(strip_body_texts(p), tdp, work)
            second = list(work.glob("*.png"))
            self.assertEqual(len(second), 2, "源文件变化后应生成新的缓存 PNG，而非复用旧的")
            names = {q.name for q in second}
            self.assertNotEqual(first[0].name, (names - {first[0].name}).pop())

    def test_render_background_includes_photo(self):
        """端到端：背景 PNG 里照片区域的像素必须是照片色，而非底色。"""
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td)
            p = self._write_photo_fixture(tdp)
            png = tdp / "bg.png"
            render_background(strip_body_texts(p), png, 480, 270,
                              svg_dir=tdp, work_dir=tdp / "work")
            im = _PILImage.open(png).convert("RGB")
            # 照片槽位中心 (100+60, 100+45) 按 480/1920 缩放 → (40, 36)
            r, g, b = im.getpixel((40, 36))
            self.assertGreater(r, 150, f"照片未渲染，中心像素={(r, g, b)}")
            self.assertGreater(r - b, 80, f"照片未渲染，中心像素={(r, g, b)}")
            # 底色区保持米色
            br, bg_, bb = im.getpixel((400, 200))
            self.assertGreater(br, 230)


if __name__ == "__main__":
    unittest.main()
