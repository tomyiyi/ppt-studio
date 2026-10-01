#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
svg_to_pptx.py —— SVG 画布 → 可编辑 PPTX 导出（对应 docs/workflow.md「阶段 8：导出并回读」）

双层结构（保真 + 可编辑）：
  背景层：去掉正文 <text> 后的 SVG 经 rsvg-convert 渲染为画布尺寸 PNG，全页铺底；
           data-decorative="true" 的装饰文字（如 0.08 透明度页码水印）保留在背景层。
  文本层：每个正文 <text> 转为真实可编辑文本框 —— 坐标 / 字号 / 颜色 / 粗细 /
           斜体 / 对齐锚点 / 字距 按 SVG 属性还原。

换算：
  字号：SVG px（96dpi）→ pt = px * 0.75（python-pptx 以 Pt 计）。
  坐标：px → 英寸按 (px / viewBox) * 幻灯片英寸 换算。
  文本框 top ≈ 基线 y - size * ASCENT_RATIO（近似上伸部，非像素级对齐）。

字号不做 spec ramp 归一 —— 忠实还原 SVG 原字号，如实交给 qa_pptx 回读裁决；
若 SVG 与 spec_lock 字阶漂移，门禁会如实报出（这是有价值的信号，不是导出器该掩盖的）。

用法：
  python3 scripts/svg_to_pptx.py projects/fw2026-trends -o output/fw2026-v4.pptx
  python3 scripts/svg_to_pptx.py projects/fw2026-trends/svg_output_v4 -o /tmp/x.pptx -f ppt169
  python3 scripts/svg_to_pptx.py --check   # 仅检查依赖（rsvg-convert / python-pptx）
"""

from __future__ import annotations

import argparse
import hashlib
import re
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from scripts.check_page_map import find_svg_dir  # 版本优先级：svg_output_v4 > svg_output

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    _PPTX_OK = True
    _PPTX_ERR = ""
except ImportError as exc:  # pragma: no cover
    _PPTX_OK = False
    _PPTX_ERR = str(exc)

try:
    from scripts.text_measure import measure as _measure_text
except Exception:  # pragma: no cover
    _measure_text = None

try:
    from PIL import Image as _PILImage
    _PIL_OK = True
except ImportError:  # pragma: no cover
    _PILImage = None
    _PIL_OK = False

SVG_NS = "http://www.w3.org/2000/svg"
XLINK_NS = "http://www.w3.org/1999/xlink"
ET.register_namespace("", SVG_NS)
# 不注册 xlink 会导致 ET 把 xlink:href 序列化成 ns1:href，rsvg 认不出 → 图片丢失
ET.register_namespace("xlink", XLINK_NS)

FORMATS = {
    "ppt169": (13.333333, 7.5),   # 16:9
    "ppt43": (10.0, 7.5),         # 4:3
}
PX_TO_PT = 0.75          # SVG px（96dpi）→ pt
ASCENT_RATIO = 0.88      # 文本框 top ≈ 基线 y - size * 0.88

if _PPTX_OK:
    _ALIGN_MAP = {"start": PP_ALIGN.LEFT, "middle": PP_ALIGN.CENTER, "end": PP_ALIGN.RIGHT}
else:
    # 无 python-pptx 时允许 --help/--check 等非构建路径正常 import；
    # 构建入口 build_pptx 会先报"缺少 python-pptx"（第 20 轮修复 NameError）。
    _ALIGN_MAP = {}
_BOLD_VALUES = {"700", "800", "900", "bold", "bolder"}


@dataclass
class TextItem:
    text: str
    x: float            # px，锚点横坐标
    y_baseline: float   # px，基线纵坐标
    size_px: float
    family: str
    fill: str           # #rrggbb
    bold: bool
    italic: bool
    anchor: str         # start | middle | end
    letter_spacing_px: float


def rsvg_available() -> bool:
    return shutil.which("rsvg-convert") is not None


def parse_viewbox(svg_path: Path) -> tuple[float, float]:
    root = ET.parse(str(svg_path)).getroot()
    vb = root.get("viewBox")
    if vb:
        parts = vb.strip().split()
        if len(parts) == 4:
            return float(parts[2]), float(parts[3])
    w = root.get("width")
    h = root.get("height")
    if w and h:
        return float(re.sub(r"[a-zA-Z%]+$", "", w)), float(re.sub(r"[a-zA-Z%]+$", "", h))
    raise ValueError(f"SVG 缺少 viewBox/width/height: {svg_path}")


def _is_decorative(el: ET.Element) -> bool:
    return el.get("data-decorative") == "true"


def extract_texts(svg_path: Path) -> list[TextItem]:
    """提取正文 <text>（跳过 data-decorative 装饰文字与空文本）。"""
    root = ET.parse(str(svg_path)).getroot()
    items: list[TextItem] = []
    for el in list(root.iter(f"{{{SVG_NS}}}text")) + list(root.iter("text")):
        if _is_decorative(el):
            continue
        text = "".join(el.itertext()).strip()
        if not text:
            continue
        try:
            x = float(el.get("x", "0"))
            y = float(el.get("y", "0"))
            size = float(el.get("font-size", "16"))
        except ValueError:
            continue
        if size <= 0:
            continue
        weight = (el.get("font-weight") or "").strip().lower()
        items.append(TextItem(
            text=text,
            x=x,
            y_baseline=y,
            size_px=size,
            family=(el.get("font-family") or "sans-serif").strip(),
            fill=(el.get("fill") or "#000000").strip(),
            bold=weight in _BOLD_VALUES,
            italic=(el.get("font-style") or "").strip().lower() == "italic",
            anchor=(el.get("text-anchor") or "start").strip().lower() or "start",
            letter_spacing_px=_parse_length(el.get("letter-spacing")),
        ))
    return items


def _parse_length(v: str | None) -> float:
    if not v:
        return 0.0
    m = re.match(r"^\s*(-?[0-9.]+)", v)
    return float(m.group(1)) if m else 0.0


def strip_body_texts(svg_path: Path) -> bytes:
    """返回去掉正文 <text> 后的 SVG 字节（装饰文字保留，用于渲染背景层）。"""
    root = ET.parse(str(svg_path)).getroot()
    for parent in list(root.iter()):
        for child in list(parent):
            if child.tag in (f"{{{SVG_NS}}}text", "text") and not _is_decorative(child):
                parent.remove(child)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


# rsvg 在本机解码 JPEG 失败（gdk-pixbuf loader 缺失，PNG 可解）→ 需要转 PNG 的后缀
_JPEG_SUFFIXES = {".jpg", ".jpeg"}

_IMAGE_HREF_RE = re.compile(
    r'(<image\b[^>]*?\b(?:href|xlink:href)\s*=\s*["\'])([^"\']+)(["\'])',
    re.IGNORECASE,
)


def _materialize_images(svg_bytes: bytes, svg_dir: Path, work_dir: Path) -> bytes:
    """把 <image> 的本地 href 改写为绝对路径；JPEG 先转 PNG 再引用。

    原因有二：(1) 背景 SVG 经 stdin 喂给 rsvg-convert，没有 base URI，
    相对路径必然解析失败；(2) 本机 rsvg 解码 JPEG 失败（loader 缺失），
    只有 PNG 能被可靠解码。data: URI 原样保留；缺失文件保留原 href
    并打印警告（rsvg 会跳过该图，不中断整页）。
    """
    svg_dir = Path(svg_dir)
    work_dir = Path(work_dir)
    work_dir.mkdir(parents=True, exist_ok=True)
    text = svg_bytes.decode("utf-8")

    def repl(m: re.Match) -> str:
        href = m.group(2).strip()
        if href.startswith("data:"):
            return m.group(0)
        src = (svg_dir / href).resolve() if not Path(href).is_absolute() else Path(href).resolve()
        if not src.is_file():
            print(f"    [warn] 背景图片缺失，rsvg 将跳过: {href}")
            return m.group(0)
        if src.suffix.lower() in _JPEG_SUFFIXES:
            if not _PIL_OK:
                raise RuntimeError(f"需要 Pillow 才能转换 JPEG 背景图: {src}")
            dst = work_dir / (hashlib.sha256(str(src).encode()).hexdigest()[:16] + ".png")
            if not dst.is_file():
                with _PILImage.open(src) as im:
                    im.convert("RGB").save(dst, "PNG")
            target = dst
        else:
            target = src
        return f"{m.group(1)}{target.as_uri()}{m.group(3)}"

    return _IMAGE_HREF_RE.sub(repl, text).encode("utf-8")


def render_background(svg_bytes: bytes, png_path: Path, width: int, height: int,
                      svg_dir: Path | None = None, work_dir: Path | None = None) -> None:
    if not rsvg_available():
        raise RuntimeError("缺少 rsvg-convert（背景层渲染依赖），请安装 librsvg")
    if svg_dir is not None:
        if work_dir is None:
            raise ValueError("传入 svg_dir 时必须同时传入 work_dir（图片物化目录）")
        svg_bytes = _materialize_images(svg_bytes, svg_dir, work_dir)
        # rsvg 从 stdin 读取 SVG 时拒绝加载外部资源（即使 file:// 绝对路径），
        # 必须落盘为文件再渲染；.bg.svg 与 .bg.png 同目录，方便 --bg-dir 调试
        staged_svg = Path(work_dir) / (Path(png_path).stem + ".bg.svg")
        staged_svg.write_bytes(svg_bytes)
        argv = ["rsvg-convert", "-w", str(width), "-h", str(height),
                "-o", str(png_path), str(staged_svg)]
        input_data = None
    else:
        argv = ["rsvg-convert", "-w", str(width), "-h", str(height), "-o", str(png_path)]
        input_data = svg_bytes
    proc = subprocess.run(argv, input=input_data, capture_output=True)
    if proc.returncode != 0 or not png_path.is_file() or png_path.stat().st_size == 0:
        raise RuntimeError(f"rsvg-convert 渲染失败: {proc.stderr.decode()[:300]}")


def parse_color(fill: str) -> tuple[int, int, int]:
    s = fill.strip().lstrip("#")
    if len(s) == 3:
        s = "".join(c * 2 for c in s)
    if len(s) >= 6:
        return int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16)
    return 0, 0, 0


def estimate_width_px(item: TextItem) -> float:
    """文本像素宽度：优先 PIL 实测，失败回退启发式；含字距补偿。"""
    w = None
    if _measure_text is not None:
        try:
            w = _measure_text(item.text, max(1, int(round(item.size_px))), item.family)
        except Exception:
            w = None
    if w is None:
        w = len(item.text) * item.size_px * 0.62
    w += item.letter_spacing_px * max(len(item.text) - 1, 0)
    return max(w + 8.0, item.size_px * 0.5)  # 8px 内边距余量


def _set_run_style(run, item: TextItem) -> None:
    run.font.size = Pt(item.size_px * PX_TO_PT)
    run.font.bold = item.bold
    run.font.italic = item.italic
    try:
        r, g, b = parse_color(item.fill)
        run.font.color.rgb = RGBColor(r, g, b)
    except Exception:
        pass
    try:
        # 东亚字体：CJK 文本走 a:ea，否则 PowerPoint 回退字体渲染
        rPr = run._r.get_or_add_rPr()
        ea = rPr.find(qn("a:ea"))
        if ea is None:
            ea = OxmlElement("a:ea")
            rPr.append(ea)
        ea.set("typeface", item.family)
        if item.letter_spacing_px:
            # spc 单位：1/100 pt
            rPr.set("spc", str(int(round(item.letter_spacing_px * PX_TO_PT * 100))))
    except Exception:
        pass


def build_pptx(
    svg_files: list[Path],
    out_path: Path,
    fmt: str = "ppt169",
    bg_dir: Path | None = None,
) -> dict:
    """核心转换：svg 文件列表 → 可编辑 pptx。返回统计。"""
    if not _PPTX_OK:  # pragma: no cover
        raise RuntimeError(f"缺少 python-pptx: {_PPTX_ERR}")
    if fmt not in FORMATS:
        raise ValueError(f"未知画幅 {fmt}，可选: {sorted(FORMATS)}")
    if not svg_files:
        raise ValueError("没有可转换的 SVG 页面")

    slide_w_in, slide_h_in = FORMATS[fmt]
    prs = Presentation()
    prs.slide_width = Inches(slide_w_in)
    prs.slide_height = Inches(slide_h_in)
    blank = prs.slide_layouts[6]

    tmp = None
    if bg_dir is None:
        tmp = tempfile.TemporaryDirectory(prefix="svg2pptx_bg_")
        bg_root = Path(tmp.name)
    else:
        bg_root = Path(bg_dir)
        bg_root.mkdir(parents=True, exist_ok=True)

    total_texts = 0
    try:
        for svg_path in svg_files:
            vb_w, vb_h = parse_viewbox(svg_path)
            texts = extract_texts(svg_path)

            bg_png = bg_root / (svg_path.stem + ".bg.png")
            render_background(
                strip_body_texts(svg_path), bg_png, int(vb_w), int(vb_h),
                svg_dir=svg_path.parent, work_dir=bg_root)

            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(
                str(bg_png), Inches(0), Inches(0),
                width=prs.slide_width, height=prs.slide_height)

            sx = slide_w_in / vb_w   # px → 英寸
            sy = slide_h_in / vb_h
            for item in texts:
                w_px = estimate_width_px(item)
                h_px = item.size_px * 1.35
                if item.anchor == "middle":
                    left_px = item.x - w_px / 2
                elif item.anchor == "end":
                    left_px = item.x - w_px
                else:
                    left_px = item.x
                top_px = item.y_baseline - item.size_px * ASCENT_RATIO

                txBox = slide.shapes.add_textbox(
                    Inches(left_px * sx), Inches(top_px * sy),
                    Inches(w_px * sx), Inches(h_px * sy))
                tf = txBox.text_frame
                tf.word_wrap = True
                for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
                    setattr(tf, m, Inches(0))
                p = tf.paragraphs[0]
                p.alignment = _ALIGN_MAP.get(item.anchor, PP_ALIGN.LEFT)
                run = p.add_run()
                run.text = item.text
                try:
                    run.font.name = item.family
                except Exception:
                    pass
                _set_run_style(run, item)
                total_texts += 1

        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out_path))
    finally:
        if tmp is not None:
            tmp.cleanup()

    return {
        "pages": len(svg_files),
        "texts": total_texts,
        "format": fmt,
        "slide_size_in": (slide_w_in, slide_h_in),
        "output": str(out_path),
    }


def resolve_svg_dir(src: str | Path | None) -> tuple[Path, Path]:
    """解析 SVG 源目录 → (svg_dir, project_dir)。"""
    if src is not None and str(src).strip() not in ("", "-"):
        p = Path(src).resolve()
        if not p.exists():
            raise FileNotFoundError(f"指定的源不存在: {src}")
        if p.is_dir() and any(p.glob("*.svg")):
            return p, p.parent
        proj = find_svg_dir(p)
        if proj is None:
            raise FileNotFoundError(f"在 {p} 下找不到 SVG 目录（svg_output*）")
        return proj, p
    # 自动发现：当前目录或 projects/ 下唯一项目（多项目 fail-closed）
    cwd = Path.cwd().resolve()
    cands: list[Path] = []
    for base in (cwd, cwd / "projects", REPO_ROOT, REPO_ROOT / "projects"):
        if not base.is_dir():
            continue
        for d in base.iterdir():
            if d.is_dir() and find_svg_dir(d) is not None:
                cands.append(d)
    uniq = sorted(set(cands))
    if len(uniq) == 1:
        proj = find_svg_dir(uniq[0])
        assert proj is not None
        return proj, uniq[0]
    if not uniq:
        raise FileNotFoundError("未发现包含 SVG 的项目目录，请显式传入项目路径")
    raise ValueError(
        "发现多个项目（%s），无法安全确定，请显式指定项目路径"
        % ", ".join(d.name for d in uniq))


def check_dependencies() -> tuple[bool, list[str]]:
    problems = []
    if not _PPTX_OK:
        problems.append(f"python-pptx 未安装: {_PPTX_ERR}")
    if not rsvg_available():
        problems.append("rsvg-convert 不可用（librsvg 未安装）")
    return (not problems, problems)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="SVG 画布 → 可编辑 PPTX 导出（背景层 + 文本层）")
    ap.add_argument("src", nargs="?", help="SVG 目录或项目目录（缺省自动发现唯一项目）")
    ap.add_argument("-o", "--output", help="输出 .pptx 路径（缺省 <项目>/output/<svg目录名>.pptx）")
    ap.add_argument("-f", "--format", default="ppt169", choices=sorted(FORMATS),
                    help="画幅（默认 ppt169 = 16:9）")
    ap.add_argument("--bg-dir", help="保留背景 PNG 的目录（默认用临时目录，转换后删除）")
    ap.add_argument("--check", action="store_true", help="仅检查依赖是否就绪")
    args = ap.parse_args(argv)

    if args.check:
        ok, problems = check_dependencies()
        for p in problems:
            print(f"  [✗] {p}")
        print("依赖就绪 ✅" if ok else "依赖缺失 ❌")
        return 0 if ok else 1

    ok, problems = check_dependencies()
    if not ok:
        for p in problems:
            print(f"[✗] {p}", file=sys.stderr)
        return 2

    try:
        svg_dir, project_dir = resolve_svg_dir(args.src)
    except (FileNotFoundError, ValueError) as exc:
        print(f"[✗] {exc}", file=sys.stderr)
        return 2

    svg_files = sorted(svg_dir.glob("*.svg"))
    out = Path(args.output).resolve() if args.output else (
        project_dir / "output" / f"{svg_dir.name}.pptx").resolve()

    try:
        stats = build_pptx(svg_files,
                           out,
                           fmt=args.format,
                           bg_dir=Path(args.bg_dir) if args.bg_dir else None)
    except Exception as exc:
        print(f"[✗] 转换失败: {exc}", file=sys.stderr)
        return 1

    print(f"已生成: {stats['output']}")
    print(f"  页面: {stats['pages']}  文本框: {stats['texts']}  画幅: {stats['format']}")
    print("下一步回读验证: python3 scripts/qa_pptx.py", stats["output"])
    return 0


if __name__ == "__main__":
    sys.exit(main())
